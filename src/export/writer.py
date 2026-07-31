"""Export the draft and its review report as DOCX / PDF (SRS FR-16).

Two artifacts, both offered from the app's Export controls:

* **the draft** — what FR-16 literally requires;
* **the review report** — conflicts (with the code-built suggested fix),
  citation warnings, template scorecard, terminology flags: the thing a desk
  officer would actually attach to the file.

Devanagari is the design constraint here. DOCX stores Unicode and Word shapes
it at render time — we just name a font with Devanagari coverage (Nirmala UI,
on every Windows government desktop). PDF writers do NO complex-script shaping
by default, which renders Marathi conjuncts visibly wrong — so the PDF path
uses fpdf2 with uharfbuzz text shaping and embedded Noto fonts (OFL — license
ships in assets/fonts/). Deterministic, zero LLM calls.
"""
import io
import re
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
FONT_DIR = ROOT / "assets" / "fonts"
sys.path.append(str(ROOT / "src" / "audit"))
from trail import _suggested_fix  # noqa: E402 — same fix text everywhere

GENERATED_BY = "ClauseGuard review report"


# ---------- shared report content ----------

def report_sections(score=None, terms=None, refs=None, review=None) -> list:
    """The report as (heading, [lines]) — one source of truth for both formats."""
    sections = []

    n_conf = ("n/a" if not review or "error" in review
              else len(review.get("conflicts", [])))
    n_cite = (len(refs.get("stale", [])) + len(refs.get("newer_related", []))
              if refs else "n/a")
    sections.append(("Review summary", [
        f"Verified conflicts: {n_conf}",
        f"Citation warnings: {n_cite}",
        f"Template compliance: {score['pct']}%" if score else "Template compliance: n/a",
        f"Terminology flags: {terms['n_flags']}" if terms else "Terminology flags: n/a",
    ]))

    if review and "error" not in review:
        lines = []
        for c in review.get("conflicts", []):
            m = c.get("meta", {})
            gr = m.get("gr_code", "?")
            lines.append(f"• Conflict with GR {gr} ({m.get('department', '?')}, "
                         f"{m.get('date', '?')})")
            lines.append(f"  Draft clause: {c.get('draft', '')}")
            lines.append(f"  Existing clause: {c.get('candidate', '')}")
            lines.append(f"  Why it conflicts: {c.get('reason', '')}")
            if fix := _suggested_fix(c.get("reason"), gr):
                lines.append(f"  Suggested fix: {fix}")
        sections.append(("Verified conflicts",
                         lines or ["No conflicts flagged against the indexed corpus."]))

    if refs:
        lines = []
        for s in refs.get("stale", []):
            lines.append(f"• Cites {s.get('cited')} which appears superseded by "
                         f"{s.get('superseded_by')}")
        for a in refs.get("newer_related", []):
            lines.append(f"• Cited GR dated {a.get('cited_date')} has "
                         f"{len(a.get('newer', []))} newer same-subject GR(s) — "
                         f"verify it is still operative")
        sections.append(("Citation warnings", lines or ["None."]))

    if score:
        fails = [c for c in score.get("checks", []) if not c["passed"]]
        lines = [f"{score['score']}/{score['max_score']} weighted points "
                 f"({score['pct']}%) across {len(score.get('checks', []))} "
                 f"Manual-of-Office-Procedure checks."]
        lines += [f"✗ {c['desc']} (weight {c['weight']})" for c in fails]
        if not fails:
            lines.append("Every template check passes.")
        sections.append(("Template compliance", lines))

    if terms:
        lines = []
        for f in terms.get("flags", []):
            approved = (f.get("approved_mr") if terms.get("primary_lang") == "mr"
                        else f.get("approved_en"))
            lines.append(f"• [{f.get('kind')}] '{f.get('found')}' → use "
                         f"'{approved}'")
        sections.append(("Terminology", lines or ["No glossary mismatches."]))

    return sections


# ---------- DOCX ----------

def _docx_base():
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Nirmala UI"          # Devanagari coverage on Windows
    style.font.size = Pt(11)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Nirmala UI")
    rfonts.set(qn("w:cs"), "Nirmala UI")    # complex-script runs (Devanagari)
    return doc


def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def draft_docx(draft_text: str) -> bytes:
    doc = _docx_base()
    for line in draft_text.splitlines():
        doc.add_paragraph(line)
    return _docx_bytes(doc)


def report_docx(score=None, terms=None, refs=None, review=None) -> bytes:
    doc = _docx_base()
    doc.add_heading(GENERATED_BY, level=1)
    doc.add_paragraph(f"Generated {datetime.now().isoformat(timespec='seconds')}")
    for heading, lines in report_sections(score, terms, refs, review):
        doc.add_heading(heading, level=2)
        for line in lines:
            doc.add_paragraph(line)
    return _docx_bytes(doc)


# ---------- PDF ----------

def _pdf_base():
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.add_font("noto", fname=str(FONT_DIR / "NotoSans-Regular.ttf"))
    pdf.add_font("notodev", fname=str(FONT_DIR / "NotoSansDevanagari-Regular.ttf"))
    pdf.set_fallback_fonts(["notodev"])     # Devanagari runs fall through here
    pdf.set_text_shaping(True)              # harfbuzz — correct conjuncts/matras
    pdf.set_auto_page_break(True, margin=18)
    pdf.set_font("noto", size=11)
    pdf.add_page()
    return pdf


def _pdf_bytes(pdf) -> bytes:
    return bytes(pdf.output())


def _pdf_text(pdf, text, size=11, pad=0.0):
    # No bold face is embedded (regular TTFs only) — emphasis is via size.
    pdf.set_font_size(size)
    pdf.multi_cell(0, size * 0.55, text)
    if pad:
        pdf.ln(pad)


def draft_pdf(draft_text: str) -> bytes:
    pdf = _pdf_base()
    _pdf_text(pdf, draft_text)
    return _pdf_bytes(pdf)


def report_pdf(score=None, terms=None, refs=None, review=None) -> bytes:
    pdf = _pdf_base()
    _pdf_text(pdf, GENERATED_BY, size=16, pad=2)
    _pdf_text(pdf, f"Generated {datetime.now().isoformat(timespec='seconds')}",
              size=9, pad=4)
    for heading, lines in report_sections(score, terms, refs, review):
        _pdf_text(pdf, heading, size=13, pad=1)
        # '✗' is in neither embedded font; keep the PDF strictly to glyphs we ship
        _pdf_text(pdf, "\n".join(l.replace("✗", "x") for l in lines), pad=3)
    return _pdf_bytes(pdf)


if __name__ == "__main__":
    # Self-test + the Devanagari gate: export Scenario A (Marathi) both ways,
    # render the PDF's first page to PNG so the conjuncts can be inspected.
    sys.stdout.reconfigure(encoding="utf-8")
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data"
    src = ROOT / "data" / "planted" / "1_fee_conflict_draft.txt"
    text = src.read_text(encoding="utf-8")

    sys.path.append(str(ROOT / "src" / "compliance"))
    from scorecard import score as score_fn
    from glossary_check import check_terminology
    s, t = score_fn(text), check_terminology(text)

    (out_dir / "export_selftest_draft.docx").write_bytes(draft_docx(text))
    (out_dir / "export_selftest_draft.pdf").write_bytes(draft_pdf(text))
    (out_dir / "export_selftest_report.docx").write_bytes(report_docx(score=s, terms=t))
    (out_dir / "export_selftest_report.pdf").write_bytes(report_pdf(score=s, terms=t))

    import fitz
    doc = fitz.open(out_dir / "export_selftest_draft.pdf")
    pix = doc[0].get_pixmap(dpi=140)
    gate = out_dir / "export_selftest_draft_page1.png"
    pix.save(gate)
    print(f"exports written to {out_dir} — INSPECT {gate.name} for conjunct shaping")
